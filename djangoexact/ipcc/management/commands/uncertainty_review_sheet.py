"""Build the review workbook for every base-value disagreement that holds an IPCC
uncertainty table back from publishing.

A disagreement means the uncertainty CSV prints one default and the reference fixture
(a snapshot of the production database) holds another, so the CSV's range cannot be
shown beside the default. Each workbook row carries what a reviewer needs to settle it
without opening another file: the IPCC source, the CSV line as printed with its range,
the database row it was joined to, and how far apart the two values are.

    python manage.py uncertainty_review_sheet
"""

import csv
from pathlib import Path

from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ipcc.management.commands import import_uncertainties as iu

DEFAULT_OUTPUT = Path("ipcc") / "data" / "uncertainties_review" / "base_mismatch_review.xlsx"

DECISIONS = {
    "keep_db": "The database value is right and the CSV range still applies to it. The CSV base is corrected to the database value; the range then publishes. Calculation results do not change.",
    "use_csv": "The CSV value is right and the database default is wrong. The database default changes, which changes results for every project using it, so it needs a pinned-number test and sign-off.",
    "drop_row": "The range cannot be trusted for this row. The CSV row is removed, so this default shows no range but the rest of the table can publish.",
    "needs_research": "Not settled yet. The table stays held back.",
}

# Only meanings stated by the cited IPCC tables; anything less certain is left blank
# rather than risk steering a reviewer with a wrong unit.
FIELD_MEANINGS = {
    ("ipcc.CropNitrousEstimationDefaultFactor", "rs_t"): "RS(T): ratio of below-ground root biomass to above-ground shoot biomass (IPCC 2019, table 11.1A)",
    ("ipcc.CropNitrousEstimationDefaultFactor", "n_bg_t"): "N_BG(T): N content of below-ground residues, kg N per kg dry matter (IPCC 2019, table 11.1A)",
    ("ipcc.ForestCombustionFactor", "co2"): "CO2 emission factor for biomass burning, g per kg dry matter burnt (IPCC 2019, table 2.5)",
    ("ipcc.ForestCombustionFactor", "ch4"): "CH4 emission factor for biomass burning, g per kg dry matter burnt (IPCC 2019, table 2.5)",
    ("ipcc.ForestCombustionFactor", "n2o"): "N2O emission factor for biomass burning, g per kg dry matter burnt (IPCC 2019, table 2.5)",
    ("ipcc.ForestCombustionFactor", "value"): "Combustion factor, fraction of biomass consumed (IPCC 2019, table 2.6)",
}

COLUMNS = (
    ("review_id", 10, "Stable number for referring to this row."),
    ("table", 30, "Reference table (Django model) the default lives in."),
    ("ipcc_source", 50, "Source cited in the uncertainty CSV header: where to look the value up."),
    ("csv_file", 30, "Uncertainty CSV under ipcc/data/uncertainties/."),
    ("csv_line", 9, "Line number in that CSV (header is line 1)."),
    ("csv_id", 9, "The CSV's own id column, when it has one."),
    ("natural_key", 60, "What the row describes, with every reference resolved to its name."),
    ("field", 14, "Column that disagrees."),
    ("field_meaning", 45, "What the column measures, where the IPCC table states it."),
    ("csv_value", 14, "Default as printed in the uncertainty CSV."),
    ("csv_min", 12, "Lower bound published with the CSV value."),
    ("csv_max", 12, "Upper bound published with the CSV value."),
    ("db_pk", 10, "Primary key of the database row the CSV row was joined to."),
    ("db_value", 14, "Default the tool calculates with today (reference fixture, refreshed from production)."),
    ("db_minus_csv", 14, "db_value minus csv_value."),
    ("diff_pct", 10, "Difference as a percentage of csv_value."),
    ("db_inside_csv_range", 14, "Whether db_value lies within [csv_min, csv_max]. 'yes' hints the range may still hold for the database value."),
    ("other_csv_columns", 45, "The CSV row's remaining columns (summary, original value, unit) as printed."),
    ("db_row", 60, "The joined database row's fields, references resolved to names, to confirm the join picked the right row."),
    ("decision", 16, "Pick one: " + ", ".join(DECISIONS) + ". See the 'How to review' sheet."),
    ("source_checked", 30, "Where you verified it (table, page, edition)."),
    ("reviewer_notes", 40, "Anything else."),
)
DECISION_COLUMN = [name for name, _width, _meaning in COLUMNS].index("decision") + 1


def read_numbered_rows(csv_path):
    """Data rows with their file line numbers, skipping blank rows as read_csv_rows does."""
    with open(csv_path, encoding=iu.csv_encoding(csv_path), newline="") as fh:
        reader = csv.reader(fh)
        next(reader)
        for row in reader:
            if any(c.strip() for c in row):
                yield reader.line_num, row


class TableContext:
    """Everything needed to describe one table's mismatching rows."""

    def __init__(self, table, fixture_cache):
        by_fixture = {s.fixture_file: s for s in iu.MANIFEST}
        self.by_model = {s.model.lower(): s for s in iu.MANIFEST}
        self.cache = fixture_cache
        self.spec = by_fixture[table["fixture"]]
        self.model = iu.django_apps.get_model(self.spec.model)
        self.csv_path = iu.data_dir() / table["csv"]
        self.field_map = {f.name: f for f in self.model._meta.concrete_fields}
        self.n2i, _rows = iu.read_csv_rows(self.csv_path)
        self.groups, ranged_names = iu.compute_ranged_groups(self.csv_path, self.n2i, self.field_map)
        colmap = iu.COLMAP.get(self.csv_path.name.lower(), {})
        self.key_cols = iu.derive_key_columns(self.n2i, colmap, self.field_map, set(self.groups), ranged_names)
        self.source_col = next((n for n in self.n2i if "source" in n.lower()), None)
        key_names = {name for name, _, _ in self.key_cols}
        self.other_cols = [n for n in self.n2i if n not in key_names | ranged_names | {"id", self.source_col}]

        fk = {mf: iu.build_fk_lookup(iu.spec_for_field(f, self.by_model), fixture_cache) for _, mf, f in self.key_cols if f.is_relation}
        fixture_rows = iu.load_fixture(self.spec, fixture_cache)
        self.db_rows = {r["pk"]: r["fields"] for r in fixture_rows}
        index = iu.build_fixture_index(fixture_rows, self.key_cols)
        self.lines_by_pk = {}
        for line_no, row in read_numbered_rows(self.csv_path):
            key = iu.build_key(row, self.key_cols, self.n2i, fk)
            pks = index.get(key) if key is not None else None
            if pks and len(pks) == 1:
                self.lines_by_pk.setdefault(pks[0], []).append((line_no, row, key))

    def names(self, field):
        rows = iu.load_fixture(iu.spec_for_field(field, self.by_model), self.cache)
        label = next((c for c in ("name", "value", "label") if rows and c in rows[0]["fields"]), None)
        return {r["pk"]: r["fields"].get(label) for r in rows} if label else {}

    def cell(self, row, ref):
        return iu.get_cell(row, ref, self.n2i).strip()

    def number(self, row, ref):
        raw = self.cell(row, ref)
        return None if iu._is_null(raw) else iu._to_float_or_none(raw)

    def csv_line(self, pk, field, csv_value):
        """The CSV line joined to `pk` whose base for `field` is the reported value."""
        base_ref = self.groups[field][0]
        for line_no, row, key in self.lines_by_pk.get(pk, []):
            if self.number(row, base_ref) == csv_value:
                return line_no, row, key
        raise LookupError(f"{self.csv_path.name}: no line joins pk {pk} with {field}={csv_value}")

    def natural_key(self, key):
        parts = []
        for (name, _model_field, field), value in zip(self.key_cols, key):
            if field.is_relation and value is not None:
                value = self.names(field).get(value, value)
            parts.append(f"{name}: {'' if value is None else value}")
        return " | ".join(parts)

    def db_row(self, pk):
        parts = []
        for name, value in self.db_rows[pk].items():
            if value is None or isinstance(value, list) or name.endswith(("_min", "_max")):
                continue
            field = self.field_map.get(name)
            if field is not None and field.is_relation:
                value = self.names(field).get(value, value)
            parts.append(f"{name}: {value}")
        return " | ".join(parts)

    def other_columns(self, row):
        return " | ".join(f"{n}: {self.cell(row, n)}" for n in self.other_cols if self.cell(row, n))

    def source(self):
        return self.source_col.split(":", 1)[-1].strip() if self.source_col else ""


ZERO_VS_EMPTY = "database empty, CSV is 0 with no width"


def inside_range(value, csv_value, lower, upper):
    if value is None:
        # Hundreds of not-applicable combinations look like this; naming them lets a
        # reviewer filter and settle them in one go instead of row by row.
        if csv_value == 0 and lower in (0, None) and upper in (0, None):
            return ZERO_VS_EMPTY
        return "database value empty"
    if lower is None or upper is None:
        return "no range published"
    return "yes" if lower <= value <= upper else "no"


def review_rows(tables, details, fixture_cache):
    """One row per disagreeing field, in COLUMNS order."""
    review_id = 0
    for table in tables:
        mismatches = details[table["csv"]].get("base_mismatch", [])
        if table["carrying_bound"] == 0 or not mismatches:
            continue
        ctx = TableContext(table, fixture_cache)
        for mismatch in mismatches:
            for field, csv_value, db_value in mismatch["f"]:
                line_no, row, key = ctx.csv_line(mismatch["pk"], field, csv_value)
                _base, min_ref, max_ref = ctx.groups[field]
                lower, upper = ctx.number(row, min_ref), ctx.number(row, max_ref)
                diff = None if db_value is None else round(db_value - csv_value, 10)
                review_id += 1
                yield [
                    review_id, table["model"], ctx.source(), table["csv"], line_no,
                    ctx.cell(row, "id") or None, ctx.natural_key(key), field,
                    FIELD_MEANINGS.get((table["model"], field), ""),
                    csv_value, lower, upper, mismatch["pk"], db_value, diff,
                    None if diff is None or csv_value == 0 else round(diff / abs(csv_value) * 100, 2),
                    inside_range(db_value, csv_value, lower, upper), ctx.other_columns(row), ctx.db_row(mismatch["pk"]),
                    None, None, None,
                ]


def write_instructions(ws):
    lines = [
        ("Why this file exists", True),
        ("Each uncertainty CSV publishes a range (min, max) for an IPCC default. A range is only shown beside a default when the CSV prints the same default the tool calculates with, so it describes the number users see.", False),
        ("The rows in 'Mismatches' are defaults where the two disagree. One disagreement holds its whole table back, so none of that table's ranges are shown. Same-number-different-precision cases (2.11 vs 2.1) already count as agreeing and are not listed.", False),
        ("", False),
        ("How to review a row", True),
        ("1. Look the value up in the IPCC source named in ipcc_source, using natural_key and field.", False),
        ("2. Compare it with csv_value and db_value. db_inside_csv_range and other_csv_columns help spot transcription slips.", False),
        ("3. Check db_row to confirm the CSV row was joined to the right database row.", False),
        ("4. Pick a decision and record where you checked in source_checked.", False),
        (f"Tip: filter db_inside_csv_range on '{ZERO_VS_EMPTY}'. Those rows have no database default and a CSV default of 0 with a 0 to 0 range, which usually means the combination does not apply; they can often be settled together.", False),
        ("", False),
        ("Decisions", True),
        *((f"{name}: {meaning}", False) for name, meaning in DECISIONS.items()),
        ("", False),
        ("Columns", True),
        *((f"{name}: {meaning}", False) for name, _width, meaning in COLUMNS),
        ("", False),
        ("Regenerate with: python manage.py uncertainty_review_sheet (from djangoexact/). Decisions typed here are not carried over, so copy them out first.", False),
    ]
    for text, heading in lines:
        ws.append([text])
        cell = ws.cell(ws.max_row, 1)
        cell.font = Font(bold=heading)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140


def write_summary(ws, tables, rows):
    range_col = [name for name, _width, _meaning in COLUMNS].index("db_inside_csv_range")
    classes = sorted({r[range_col] for r in rows})
    ws.append(["table", "csv_file", "ipcc_source", "fields_to_review", "ranges_held_back", *classes])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for table in tables:
        table_rows = [r for r in rows if r[1] == table["model"]]
        if table_rows:
            per_class = [sum(1 for r in table_rows if r[range_col] == c) for c in classes]
            ws.append([table["model"], table["csv"], table_rows[0][2], len(table_rows), table["carrying_bound"], *per_class])
    for i, width in enumerate((40, 45, 90, 16, 18, *(18 for _ in classes)), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def write_mismatches(ws, rows):
    ws.append([name for name, _width, _meaning in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, (_name, width, _meaning) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = ws.dimensions

    if ws.max_row > 1:
        letter = get_column_letter(DECISION_COLUMN)
        choices = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
        ws.add_data_validation(choices)
        choices.add(f"{letter}2:{letter}{ws.max_row}")
        fill = PatternFill("solid", fgColor="FFF2CC")
        for (cell,) in ws.iter_rows(min_row=2, min_col=DECISION_COLUMN, max_col=DECISION_COLUMN):
            cell.fill = fill


class Command(BaseCommand):
    help = "Write the review workbook for every base-value disagreement holding an IPCC uncertainty table back."

    def add_arguments(self, parser):
        parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Workbook path, relative to djangoexact/.")

    def handle(self, *args, **options):
        fixture_cache, details = {}, {}
        cap = iu.DETAIL_CAP
        iu.DETAIL_CAP = 10**9  # The committed detail report is capped; a review needs every row.
        try:
            tables, _results = iu.compute_all(sorted(iu.data_dir().glob("*.csv")), fixture_cache, details)
        finally:
            iu.DETAIL_CAP = cap

        rows = list(review_rows(tables, details, fixture_cache))
        wb = Workbook()
        wb.active.title = "How to review"
        write_instructions(wb.active)
        write_summary(wb.create_sheet("Summary"), tables, rows)
        write_mismatches(wb.create_sheet("Mismatches"), rows)

        output = Path(options["output"])
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)
        self.stdout.write(self.style.SUCCESS(f"{len(rows)} disagreements written to {output}"))
