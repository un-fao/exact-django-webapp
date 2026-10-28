"""Excel Results rows for land modules: one row per carbon pool, each tonne once.

DB-free: the report classes are built without running a calculator and fed
emission lines in the cached-results format.

Run with:
    python manage.py test api.tests.reports.test_land_pool_rows
"""
from __future__ import annotations

import unittest
from types import SimpleNamespace

import openpyxl

import math_model.no_time_dependency_final.ghg_emissions_classes as math_utils
from api.reports import land_pools
from api.reports.base import BaseProjectReport
from api.reports.constants import ROW_CUMULATIVE_HECTARES, ROW_FIRST_CATEGORY, ROW_YEARLY_BALANCE, SUMMARY_CATEGORY_ROWS
from api.reports.data_types import ActivityResult
from api.reports.extractors import extract_emissions
from api.reports.land import (
    CoastalWetlandReport,
    ForestManagementReport,
    GrasslandReport,
    LandUseChangeReport,
)
from api.reports.renderer import ExcelRenderer

AT = math_utils.ActivityTypes
GT = math_utils.GasTypes
DURATION = 3


def _balance(lines: dict[tuple, float]) -> list[dict]:
    """Cached-format emission lines; year i of a line is ``value * (i + 1)``."""
    return [
        {
            "activity": activity.value,
            "gas_type": gas.name,
            "emissions": [{"value": value * (i + 1)} for i in range(DURATION)],
        }
        for (activity, gas), value in lines.items()
    ]


def _yearly(value: float) -> list[float]:
    return [value * (i + 1) for i in range(DURATION)]


def _report(cls, lines: dict[tuple, float]):
    report = cls.__new__(cls)
    report.emissions_set = _balance(lines)
    report.activity_report = SimpleNamespace(project_report=SimpleNamespace(duration=DURATION))
    return report


def _rows(report, soil_fire=None):
    soil_co2, soil_n2o, fire_n2o, fire_ch4 = soil_fire or report._extract_land_base()
    return {r.label: r.yearly_values for r in report._standard_result_rows(soil_co2, soil_n2o, fire_n2o, fire_ch4)}


def _row_sum(rows: dict) -> list[float]:
    return [sum(values[i] for values in rows.values()) for i in range(DURATION)]


FOREST_LINES = {
    (AT.AGB_GROWTH, GT.CO2): -100.0,
    (AT.BGB_GROWTH, GT.CO2): -25.0,
    (AT.ROTATION_AGB, GT.CO2): 4.0,
    (AT.ROTATION_AGB, GT.N2O): 0.1,
    (AT.ROTATION_AGB, GT.CH4): 0.2,
    (AT.ROTATION_BGB, GT.CO2): 1.0,
    (AT.HWP_ROTATION_AGB, GT.CO2): 30.0,
    (AT.HWP_ROTATION_BGB, GT.CO2): 7.0,
    (AT.LOGGING_AGB, GT.CO2): 2.0,
    (AT.HWP_LOGGING_AGB, GT.CO2): 3.0,
    (AT.DISTURBANCE_AGB, GT.CO2): 5.0,
    (AT.DISTURBANCE_FIRE_AGB, GT.CO2): 6.0,
    (AT.DISTURBANCE_FIRE_AGB, GT.N2O): 0.3,
    (AT.DISTURBANCE_FIRE_BGB, GT.CO2): 1.5,
    (AT.DEGRADATION_AGB, GT.CO2): 8.0,
    (AT.DEGRADATION_BGB, GT.CO2): 2.0,
    (AT.LITTER, GT.CO2): -4.0,
    (AT.DEADWOOD, GT.CO2): -3.0,
    (AT.DEGRADATION_LITTER, GT.CO2): 0.5,
    (AT.DEGRADATION_DEADWOOD, GT.CO2): 0.25,
    (AT.SOIL_CO2_CHANGE, GT.CO2): -11.0,
    (AT.SOM, GT.N2O): 0.7,
}


class TestForestManagementRows(unittest.TestCase):
    def setUp(self):
        report = _report(ForestManagementReport, FOREST_LINES)
        soil_co2, soil_n2o, _, _ = report._extract_land_base()
        # Forest fire N2O / CH4 come from the rotation, logging and disturbance-fire lines.
        self.report = report
        self.rows = _rows(report, (soil_co2, soil_n2o, _yearly(0.1 + 0.3), _yearly(0.2)))

    def test_one_row_per_pool(self):
        self.assertEqual(list(self.rows), [
            land_pools.LABEL_AGB,
            land_pools.LABEL_BGB,
            land_pools.LABEL_DOM,
            "CO2 in soils",
            "N2O in soils",
            "N2O from fires",
            "CH4 from fires",
        ])

    def test_pool_values(self):
        expected = {
            land_pools.LABEL_AGB: -100 + 4 + 30 + 2 + 3 + 5 + 6 + 8,
            land_pools.LABEL_BGB: -25 + 1 + 7 + 1.5 + 2,
            land_pools.LABEL_DOM: -4 - 3 + 0.5 + 0.25,
            "CO2 in soils": -11.0,
            "N2O in soils": 0.7,
        }
        for label, value in expected.items():
            for got, want in zip(self.rows[label], _yearly(value)):
                self.assertAlmostEqual(got, want, places=9, msg=label)

    def test_rows_add_up_to_module_total(self):
        for got, want in zip(_row_sum(self.rows), self.report._balance_total()):
            self.assertAlmostEqual(got, want, places=9)


class TestLandUseChangeRows(unittest.TestCase):
    LINES = {
        (AT.BIOMASS, GT.CO2): 500.0,
        (AT.DOM, GT.CO2): 40.0,
        (AT.SOIL_CO2_CHANGE, GT.CO2): 60.0,
        (AT.SOM, GT.N2O): 2.0,
        (AT.RESIDUE_BURNING, GT.N2O): 1.0,
        (AT.RESIDUE_BURNING, GT.CH4): 3.0,
    }

    def test_dom_has_its_own_row_and_leaves_soils(self):
        report = _report(LandUseChangeReport, self.LINES)
        rows = _rows(report)
        self.assertEqual(list(rows)[:3], [land_pools.LABEL_BIOMASS, land_pools.LABEL_DOM, "CO2 in soils"])
        self.assertEqual(rows[land_pools.LABEL_DOM], _yearly(40.0))
        self.assertEqual(rows["CO2 in soils"], _yearly(60.0))
        self.assertEqual(_row_sum(rows), report._balance_total())


class TestUnsplitModules(unittest.TestCase):
    def test_grassland_keeps_one_biomass_row(self):
        rows = _rows(_report(GrasslandReport, {(AT.BIOMASS, GT.CO2): -9.0, (AT.SOIL_CO2_CHANGE, GT.CO2): -1.0}))
        self.assertIn(land_pools.LABEL_BIOMASS, rows)
        self.assertNotIn(land_pools.LABEL_AGB, rows)
        self.assertNotIn("Other emissions (CO2-eq)", rows)

    def test_coastal_wetland_names_dom_inside_its_biomass_row(self):
        rows = _rows(_report(CoastalWetlandReport, {(AT.BIOMASS, GT.CO2): 12.0}), soil_fire=[[0.0] * DURATION] * 4)
        self.assertEqual(rows["CO2 in biomass and dead organic matter (AGB+BGB+litter+deadwood)"], _yearly(12.0))

    def test_lines_without_a_row_land_in_other_emissions(self):
        report = _report(GrasslandReport, {
            (AT.BIOMASS, GT.CO2): -9.0,
            (AT.REWETTING_REVEGETATION, GT.CH4): 4.0,
        })
        rows = _rows(report)
        self.assertEqual(rows["Other emissions (CO2-eq)"], _yearly(4.0))
        self.assertEqual(_row_sum(rows), report._balance_total())


class TestProjectSummary(unittest.TestCase):
    """Rows 4 onwards of the Results sheet: pools split, every category once."""

    def setUp(self):
        luc_lines = TestLandUseChangeRows.LINES
        reports = []
        for lines in (FOREST_LINES, luc_lines, {(AT.FUEL, GT.CO2): 10.0, (AT.CH4_EMITTED_RICE, GT.CH4): 2.0}):
            es = _balance(lines)
            reports.append(SimpleNamespace(emissions_set=es, emissions_set_w=es, emissions_set_wo=[]))
        self.all_lines = [line for r in reports for line in r.emissions_set]

        project_report = BaseProjectReport.__new__(BaseProjectReport)
        project_report.project = SimpleNamespace(implementation_years=DURATION, capitalization_years=0)
        activity = ActivityResult("A", [], [], [], _module_reports=reports)
        self.agg = project_report._build_aggregated_from_module_emissions([activity])

    def _expected(self, activities):
        total = [0.0] * DURATION
        for a in activities:
            values = extract_emissions(self.all_lines, a, GT.CO2, duration=DURATION)
            total = [t + v for t, v in zip(total, values)]
        return total

    def test_categories_add_up_to_every_emission_line(self):
        everything = extract_emissions(self.all_lines, duration=DURATION)
        for got, want in zip(self.agg.yearly_balance, everything):
            self.assertAlmostEqual(got, want, places=9)
        for got, want in zip(self.agg.yearly_balance_w, everything):
            self.assertAlmostEqual(got, want, places=9)

    def test_forest_pools_are_not_other_co2(self):
        self.assertEqual(self.agg.agb_co2, self._expected(land_pools.AGB))
        self.assertEqual(self.agg.bgb_co2, self._expected(land_pools.BGB))
        self.assertEqual(self.agg.dom_co2, self._expected(land_pools.DOM))
        self.assertEqual(self.agg.biomass_co2, _yearly(500.0))
        self.assertEqual(self.agg.other_co2, _yearly(10.0))

    def test_results_sheet_rows(self):
        renderer = ExcelRenderer.__new__(ExcelRenderer)
        renderer.result = SimpleNamespace(
            aggregated=self.agg, start_year=2020, last_year=2020 + DURATION,
            cumulative_hectares_yearly=[], cumulative_heads_yearly=[], cumulative_catch_yearly=[],
        )
        wb = openpyxl.Workbook()
        wb.active.title = "Results"
        renderer._write_results_summary(wb)
        ws = wb["Results"]
        for offset, (_, attr) in enumerate(SUMMARY_CATEGORY_ROWS):
            self.assertEqual(ws.cell(ROW_FIRST_CATEGORY + offset, 2).value, getattr(self.agg, attr)[0])
        category_sum = sum(ws.cell(row, 2).value for row in range(ROW_FIRST_CATEGORY, ROW_CUMULATIVE_HECTARES))
        self.assertAlmostEqual(category_sum, ws.cell(ROW_YEARLY_BALANCE, 2).value, places=9)


if __name__ == "__main__":
    unittest.main()
