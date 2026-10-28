"""Project indicators: registry, land groups on real models, endpoint, Excel sheet.

Run with:
    python manage.py test api.tests.test_indicators
"""

from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import Group
from django.test import SimpleTestCase, TestCase
from rest_framework.test import APIRequestFactory, force_authenticate

from api.models import AnnualCropland, Country, ModuleType, ProjectMembership
from api.reports.data_types import ActivityResult
from api.reports.renderer import ExcelRenderer
from api.services import indicators
from api.tests.unit import base_module
from api.tests.unit.factories import UnitTestAnnualCroplandFactory
from api.tests.factories import (
    ActivityFactory,
    AnnualCroplandFactory,
    ForestManagementFactory,
    GrasslandFactory,
    LandUseChangeFactory,
    ProjectFactory,
    UserFactory,
)
from api.views import ProjectViewSet

LAND_GROUPS = ["land_use_change_matrix"]
# Dropped after the 2026-10-28 methodology review; they must not come back unnoticed.
REMOVED_GROUPS = {"soil_carbon", "four_per_1000", "forest_cover", "net_zero_deforestation"}


class RegistryTests(SimpleTestCase):
    def tearDown(self):
        indicators._GROUPS.pop("test_extra", None)
        indicators._GROUPS.pop("test_broken", None)

    def test_registered_group_is_built_without_touching_the_core(self):
        @indicators.register_group("test_extra", "Extra", order=999)
        def extra(ctx):
            return {"indicators": [indicators.indicator("x", "X", "t", value=len(ctx.activity_results))]}

        project = SimpleNamespace(pk=1, implementation_years=5, capitalization_years=15, start_year_of_activities=2025)
        out = indicators.build(project, [])
        ids = [g["id"] for g in out["groups"]]
        self.assertEqual(ids[: len(LAND_GROUPS)], LAND_GROUPS)
        self.assertEqual(ids[-1], "test_extra")
        self.assertEqual(out["groups"][-1]["indicators"][0]["value"], 0)
        self.assertTrue(out["groups"][-1]["available"])

    def test_failing_group_is_unavailable_not_fatal(self):
        @indicators.register_group("test_broken", "Broken")
        def broken(ctx):
            raise RuntimeError("boom")

        project = SimpleNamespace(pk=1, implementation_years=5, capitalization_years=15, start_year_of_activities=2025)
        with patch.object(indicators.logger, "exception") as log_exception:
            out = indicators.build(project, [])
        log_exception.assert_called_once()
        group = next(g for g in out["groups"] if g["id"] == "test_broken")
        self.assertFalse(group["available"])

    def test_project_without_land_has_every_land_group_unavailable(self):
        project = SimpleNamespace(pk=1, implementation_years=5, capitalization_years=15, start_year_of_activities=2025)
        out = indicators.build(project, [])
        self.assertFalse(any(g["available"] for g in out["groups"] if g["id"] in LAND_GROUPS))


def _module_type(class_name):
    return ModuleType.objects.get(class_name=class_name)


class LandGroupsTests(TestCase):
    def setUp(self):
        self.user = UserFactory(email="indicators@example.com")
        self.project = ProjectFactory(
            owner=self.user, country=Country.objects.first(), implementation_years=6, start_year_of_activities=2024,
        )
        ProjectMembership.objects.create(user=self.user, project=self.project, group=Group.objects.get(name="Admin"))

        self.deforestation = ActivityFactory(project=self.project)
        forest, cropland = _module_type("ForestManagement"), _module_type("AnnualCropland")
        self.deforestation.module_types.add(forest, cropland, _module_type("LandUseChange"))
        luc = LandUseChangeFactory(
            activity=self.deforestation, area=10,
            module_type_start=forest, module_type_w=cropland, module_type_wo=forest,
        )
        ForestManagementFactory(activity=self.deforestation, area=10, land_use_change=luc)
        AnnualCroplandFactory(activity=self.deforestation, area=10, land_use_change=luc)

        self.pasture = ActivityFactory(project=self.project)
        self.pasture.module_types.add(_module_type("Grassland"))
        GrasslandFactory(activity=self.pasture, area=5)

    def _build(self):
        results = [
            ActivityResult(title=a.name, module_results=[], total_emissions=[], total_hectares_yearly=[], _activity=a)
            for a in (self.deforestation, self.pasture)
        ]
        return {g["id"]: g for g in indicators.build(self.project, results)["groups"]}

    @staticmethod
    def _values(group):
        return {i["id"]: i for i in group["indicators"]}

    def test_matrix_from_luc_and_unchanged_land(self):
        group = self._build()["land_use_change_matrix"]
        cats = [c["id"] for c in group["data"]["categories"]]
        forest, cropland, grassland = cats.index("forest"), cats.index("cropland"), cats.index("grassland")
        self.assertEqual(group["data"]["with"][forest][cropland], 10)
        self.assertEqual(group["data"]["without"][forest][forest], 10)
        self.assertEqual(group["data"]["with"][grassland][grassland], 5)
        converted = self._values(group)["converted_area"]
        self.assertEqual((converted["with"], converted["without"]), (10, 0))
        self.assertEqual(set(converted), {"id", "label", "unit", "with", "without", "value", "series"})

    def test_endpoint_permissions(self):
        view = ProjectViewSet.as_view({"get": "indicators"})
        factory = APIRequestFactory()

        request = factory.get(f"/api/projects/{self.project.pk}/indicators/", {"lang": "fr"})
        force_authenticate(request, user=self.user)
        response = view(request, pk=self.project.pk)
        self.assertEqual(response.status_code, 200, response.data)
        ids = [g["id"] for g in response.data["groups"]]
        self.assertEqual(ids[: len(LAND_GROUPS)], LAND_GROUPS)
        self.assertFalse(REMOVED_GROUPS & set(ids))

        stranger = UserFactory(email="indicators-stranger@example.com")
        request = factory.get(f"/api/projects/{self.project.pk}/indicators/")
        force_authenticate(request, user=stranger)
        self.assertEqual(view(request, pk=self.project.pk).status_code, 403)


class ExcelSheetTests(SimpleTestCase):
    EXISTING_SHEETS = ["Results", "Metadata", "Additional Indicators", "Shadow Price of Carbon", "Inventory"]

    def test_indicators_sheet_is_appended_after_existing_sheets(self):
        data = {"groups": [
            {"id": "land_use_change_matrix", "label": "Land use change matrix", "available": True, "data": None,
             "indicators": [indicators.indicator("converted_area", "Area changing land use", "ha", with_=10, without=0)]},
            {"id": "food_waste", "label": "Food waste", "available": False, "data": None, "indicators": []},
        ]}
        renderer = ExcelRenderer(SimpleNamespace(indicators=data))
        wb = renderer._manager.get_workbook()
        renderer._write_indicators(wb)
        self.assertEqual(wb.sheetnames, self.EXISTING_SHEETS + ["Indicators"])
        ws = wb["Indicators"]
        self.assertEqual([ws.cell(row=2, column=c).value for c in range(1, 6)],
                         ["Indicator", "Unit", "With project", "Without project", "Value"])
        self.assertEqual(ws.cell(row=3, column=1).value, "Area changing land use")
        self.assertEqual(ws.cell(row=3, column=3).value, 10)

    def test_no_indicators_leaves_workbook_unchanged(self):
        renderer = ExcelRenderer(SimpleNamespace(indicators=None))
        wb = renderer._manager.get_workbook()
        renderer._write_indicators(wb)
        self.assertEqual(wb.sheetnames, self.EXISTING_SHEETS)


class ComputedProjectTests(base_module.BaseModuleTestCase):
    """A real, calculated annual cropland project: indicators read its results and change none."""

    test_defaults_visible_with_different_lock_holder = None
    test_land_use_change_calculation = None
    test_module_edit_invalidates_cache = None
    test_generic_calculation = None

    def setUp(self):
        self.ModuleClass = AnnualCropland
        super().setUp()
        self.edit_module(self.module, self.user, UnitTestAnnualCroplandFactory.get_validated_data())
        self.module.refresh_from_db()

    def test_indicators_read_results_without_changing_them(self):
        before = self.get_results(cached="false").data["balance"]
        out = indicators.compute_project_indicators(self.project)
        after = self.get_results(cached="false").data["balance"]
        self.assertEqual(before, after)

        self.assertEqual(out["skipped_activities"], [])
        groups = {g["id"]: g for g in out["groups"]}
        matrix = groups["land_use_change_matrix"]["data"]["with"]
        self.assertEqual(matrix[1][1], self.module.area)
        self.assertTrue(groups["land_use_change_matrix"]["available"])
        self.assertFalse(REMOVED_GROUPS & set(groups))

    def test_excel_report_gains_an_indicators_sheet(self):
        from io import BytesIO

        import openpyxl

        from api.reports import generate_excel_report

        wb = openpyxl.load_workbook(BytesIO(generate_excel_report(self.project)))
        self.assertEqual(wb.sheetnames[-1], "Indicators")


class PdfPartialTests(SimpleTestCase):
    def test_partial_renders_available_groups_and_matrix(self):
        from django.template.loader import render_to_string

        data = {"groups": [
            {"id": "land_use_change_matrix", "label": "Land use change matrix", "available": True,
             "indicators": [indicators.indicator("converted_area", "Area changing land use", "ha", with_=12.5, without=0)],
             "data": {"categories": [{"id": "forest", "label": "Forest"}, {"id": "cropland", "label": "Cropland"}],
                      "with": [[0, 12.5], [0, 3]], "without": [[12.5, 0], [0, 3]]}},
            {"id": "food_waste", "label": "Hidden group", "available": False, "indicators": [], "data": None},
        ]}
        html = render_to_string("reports/partials/common/indicators.html", {"indicators": data})
        self.assertIn("Area changing land use", html)
        self.assertIn("12.50", html)
        self.assertIn("<td>Cropland</td>", html)
        self.assertNotIn("Hidden group", html)
