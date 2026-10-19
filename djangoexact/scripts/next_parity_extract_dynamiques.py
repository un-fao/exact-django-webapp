"""Regenerate the NEXT dynamics golden file from NEXT's hidden `Dynamiques` sheet.

The repo never reads the NEXT workbook at test time. Run this by hand, in a
throwaway venv (not the repo one), whenever the goldens need refreshing:

    python -m venv /tmp/next-venv && /tmp/next-venv/bin/pip install openpyxl formulas
    /tmp/next-venv/bin/python -I djangoexact/scripts/next_parity_extract_dynamiques.py \\
        path/to/NEXT062026_unlocked.xlsx \\
        djangoexact/math_model/tests/goldens/next_dynamiques.json

Each series cell's Excel formula is compiled with the `formulas` package and
evaluated with the scenario's inputs bound to NEXT's named ranges. Before
writing anything the script checks that this reproduces the values Excel
cached for the sheet's own inputs.
"""

import argparse
import json
import math

import formulas
import openpyxl

SHEET = "Dynamiques"
INPUT_CELLS = {"time": "H3", "time2": "H4", "x0": "O3", "x_without": "O4", "x_with": "O5"}
# Excel named range -> scenario key. time_tot (H5) is itself a formula, see evaluate_cell.
NAMED_RANGES = {"TIME": "time", "TIME2": "time2", "X0": "x0", "XNO": "x_without", "XYES": "x_with"}
SERIES_COLUMNS = {
    "linear": {"without": "D", "with": "F"},
    "logistic": {"without": "S", "with": "T"},
}
T_COLUMNS = ("C", "R")
T0_ROW = 9
# Dynamiques only has rows for t = 0..20; past time_tot every cell is #N/A.
MAX_T = 20

SCENARIOS = [
    {"id": "next_default", "time": 5, "time2": 7, "x0": 2000, "x_without": 1000, "x_with": 7000},
    {"id": "adoption_0_to_100_over_20", "time": 20, "time2": 0, "x0": 0, "x_without": 0, "x_with": 100},
    {"id": "decrease_odd_time", "time": 3, "time2": 10, "x0": 50, "x_without": 40, "x_with": 10.5},
    {"id": "exact_default_4_16", "time": 4, "time2": 16, "x0": 12.5, "x_without": 12.5, "x_with": 30},
]

METADATA = {
    "source": "NEXT062026_unlocked.xlsx, hidden sheet Dynamiques (rows 9-29, columns D/F linear, S/T logistic)",
    "inputs": {
        "time": "implementation length in years (named range time, Dynamiques!H3)",
        "time2": "capitalisation length in years (time2, H4); time_tot = time + time2",
        "x0": "value at t=0 (X0, O3)",
        "x_without": "end value of the without-project series (Xno, O4)",
        "x_with": "end value of the with-project series (Xyes, O5)",
    },
    "series": (
        "Each series has one point per integer t = 0..time_tot (time_tot + 1 points). "
        "Point t is the state reached after t years, no half-year averaging. For t >= time the value is the end value."
    ),
    "formulas": {
        "linear": "t >= time ? X1 : X0 + t * (X1 - X0) / time",
        "logistic": (
            "t >= time ? X1 : X0 + (X1 - X0) * (1 / (1 + exp(-(10 / time) * (t - time / 2))) - 1 / (1 + exp(5)))"
        ),
        "logistic_notes": (
            "Steepness k = 10 / time, midpoint time / 2. The 1 / (1 + exp(5)) offset makes X(0) == X0 exactly. "
            "The curve only reaches about 98.66% of (X1 - X0) just before t = time, then jumps to X1 at t = time."
        ),
    },
    "calc_sheet_usage": (
        "Calculation sheets (e.g. Ag land row 26) inline the same formulas, picking L/S/E from a selector cell, "
        "with time = end_year - start_year and t = year - start_year + 1 (the start year already gets t = 1). "
        "Exponential (E) is not exported here."
    ),
    "exact_comparison": (
        "EX-ACT compute_yearly_or_half_year_cumulative(start, end, years_implementation, years_capitalization, fn, "
        "interim_values=False) returns the same time_tot + 1 points. With interim_values=True (the default) it "
        "returns average_yearly_value(points), the mean of consecutive points (mid-year values, time_tot items). "
        "Compare like with like: interim=False against the points, or interim=True against "
        "average_yearly_value(points)."
    ),
}


def evaluate_cell(formula: str, row: int, inputs: dict, time_tot_formula: str) -> float:
    func = formulas.Parser().ast(formula)[1].compile()
    args = []
    for key in func.inputs:
        bare = key.replace("$", "")
        if bare in NAMED_RANGES:
            args.append(inputs[NAMED_RANGES[bare]])
        elif bare == "TIME_TOT":
            args.append(evaluate_cell(time_tot_formula, row, inputs, time_tot_formula))
        elif bare == f"{bare[0]}{row}" and bare[0] in T_COLUMNS:
            args.append(row - T0_ROW)
        else:
            raise ValueError(f"Unexpected precedent {key} in {formula}")
    return float(func(*args))


def evaluate(ws, inputs: dict) -> dict:
    time_tot = inputs["time"] + inputs["time2"]
    return {
        fn: {
            side: [
                evaluate_cell(ws[f"{col}{T0_ROW + t}"].value, T0_ROW + t, inputs, ws["H5"].value)
                for t in range(time_tot + 1)
            ]
            for side, col in cols.items()
        }
        for fn, cols in SERIES_COLUMNS.items()
    }


def check_against_cache(ws, cached_ws) -> None:
    inputs = {key: cached_ws[cell].value for key, cell in INPUT_CELLS.items()}
    for fn, sides in evaluate(ws, inputs).items():
        for side, values in sides.items():
            col = SERIES_COLUMNS[fn][side]
            expected = [cached_ws[f"{col}{T0_ROW + t}"].value for t in range(len(values))]
            assert all(math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-9) for a, b in zip(values, expected)), (
                f"formulas disagrees with Excel's cached {fn}/{side}: {values} != {expected}"
            )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("output")
    args = parser.parse_args()

    ws = openpyxl.load_workbook(args.workbook, keep_links=False)[SHEET]
    check_against_cache(ws, openpyxl.load_workbook(args.workbook, data_only=True, keep_links=False)[SHEET])

    scenarios = []
    for scenario in SCENARIOS:
        inputs = {key: value for key, value in scenario.items() if key != "id"}
        assert inputs["time"] + inputs["time2"] <= MAX_T
        series = evaluate(ws, inputs)
        t = list(range(len(series["linear"]["with"])))
        scenarios.append({"id": scenario["id"], "inputs": inputs, "t": t, **series})

    metadata = dict(METADATA, formulas=dict(METADATA["formulas"]))
    metadata["formulas"]["excel_linear_with"] = ws[f"F{T0_ROW}"].value
    metadata["formulas"]["excel_logistic_with"] = ws[f"T{T0_ROW}"].value
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump({"metadata": metadata, "scenarios": scenarios}, f, indent=2)
        f.write("\n")


if __name__ == "__main__":
    main()
